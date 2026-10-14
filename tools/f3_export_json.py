#!/usr/bin/env python3
"""Fase 3/5 (WS-SUPABASE): llig l'export .xlsx del Google Sheet i el torna per stdout com a JSON.

Només LLIG. Cada pestanya passa a una llista de files {capçalera: text}. L'usa tools/supabase_carrega.mjs.
Ús: python -I tools/f3_export_json.py "C:/ruta/Falla Portal DB.xlsx"
"""
import datetime as dt
import json
import sys

from openpyxl import load_workbook


def text(v):
    if v is None:
        return ''
    if isinstance(v, bool):
        return 'true' if v else 'false'
    if isinstance(v, dt.datetime):
        return v.date().isoformat() if v.time() == dt.time(0, 0) else v.isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, float):
        return str(int(v)) if v == int(v) else repr(v)
    return str(v)


def main():
    if len(sys.argv) != 2:
        sys.exit('Ús: python -I tools/f3_export_json.py <export.xlsx>')
    wb = load_workbook(sys.argv[1], read_only=True, data_only=True)
    out = {}
    for ws in wb.worksheets:
        it = ws.iter_rows(values_only=True)
        try:
            capcalera = [text(c) for c in next(it)]
        except StopIteration:
            out[ws.title] = []
            continue
        files = []
        for r in it:
            if any(c not in (None, '') for c in r):
                files.append({h: text(c) for h, c in zip(capcalera, r) if h})
        out[ws.title] = files
    sys.stdout.reconfigure(encoding='utf-8')
    json.dump(out, sys.stdout, ensure_ascii=False)


main()
