#!/usr/bin/env python3
"""Fase 0 (WS-SUPABASE): perfil de les dades reals del Google Sheet de Falla Portal.

Només LLIG l'export .xlsx; no modifica res i no es connecta a cap servei.

Ús (des de la carpeta del projecte):
    python -I tools/f0_perfil_dades.py "C:/ruta/FallaPortal_export.xlsx"

Com obtenir l'export: al Google Sheet, Fitxer -> Descarrega -> Microsoft Excel (.xlsx).
Guarda'l fora del repositori o amb un nom que comence per `FallaPortal_export`
(el .gitignore l'exclou). L'informe es desa al costat de l'export com
`<export>_perfil_dades_<data>.md` (també ignorat per git) i conté xifres reals: no el pugis enlloc.

Comprovacions: vegeu INFORME_ANOMALIES_F0.md §5.
"""
import re
import sys
import datetime as dt
from collections import Counter, defaultdict

from openpyxl import load_workbook

TABS = {
    'moviments': ('01 Moviments', ['id', 'data', 'tipus', 'via', 'areaId', 'areaNom', 'delegacioId', 'delegacioNom', 'eventId', 'eventNom', 'concepte', 'establiment', 'responsable', 'responsableEmail', 'import', 'notes', 'anyFiscal', 'origenFons', 'fontImport', 'bankHash', 'origen', 'ticketUrl', 'ticketMotiu', 'conciliacio']),
    'fons': ('02 Fons', None),
    'tancaments': ('03 Tancaments', None),
    'arees': ('04 Arees', None),
    'delegacions': ('05 Delegacions', None),
    'events': ('06 Events', None),
    'bancs': ('07 Bancs', None),
    'usuaris': ('08 Usuaris', None),
    'assignacions': ('09 Usuaris Assignacions', None),
    'pres_plans': ('10 Pressupostos Plans', None),
    'pres_config': ('11 Pressupostos Config', None),
    'pres_historial': ('12 Pressupostos Historial', None),
    'regles': ('13 Regles CSV', None),
    'rols': ('14 Rols', None),
    'saldos': ('15 Saldos', None),
    'fapp_passiu': ('16 Fapp Passiu', None),
    'config': ('00 Config', None),
}

RE_DATA = re.compile(r'^\d{2}/\d{2}/\d{4}$')
RE_ISO = re.compile(r'^\d{4}-\d{2}-\d{2}')
RE_ANY = re.compile(r'^\d{2}/\d{2}$')
OUT = []


def out(s=''):
    OUT.append(s)
    print(s)


def cell_str(v):
    if v is None:
        return ''
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime('%d/%m/%Y')
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def read_tabs(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    data = {}
    missing = []
    for key, (name, _) in TABS.items():
        if name not in wb.sheetnames:
            missing.append(name)
            data[key] = []
            continue
        rows = list(wb[name].iter_rows(values_only=True))
        if len(rows) < 2:
            data[key] = []
            continue
        head = [cell_str(h) for h in rows[0]]
        recs = []
        for r in rows[1:]:
            if not any(c not in (None, '') for c in r):
                continue
            recs.append({h: cell_str(r[i]) if i < len(r) else '' for i, h in enumerate(head) if h})
        data[key] = recs
    return data, missing, wb.sheetnames


def num(s):
    """Interpreta un import tal com el guarda l'app (text amb punt decimal o coma)."""
    s = (s or '').replace(' ', '').replace('€', '')
    if not s:
        return None
    try:
        return float(s)
    except ValueError:
        pass
    if ',' in s and s.count(',') == 1:
        try:
            return float(s.replace('.', '').replace(',', '.'))
        except ValueError:
            return None
    return None


def section(t):
    out()
    out('## ' + t)
    out()


def table(rows, header):
    out('| ' + ' | '.join(header) + ' |')
    out('|' + '---|' * len(header))
    for r in rows:
        out('| ' + ' | '.join(str(x) for x in r) + ' |')
    out()


def empties(recs, cols):
    n = len(recs)
    res = []
    for c in cols:
        b = sum(1 for r in recs if not r.get(c))
        if b:
            res.append((c, b, f'{100 * b / n:.0f}%'))
    return res


def main(path):
    d, missing, names = read_tabs(path)
    out(f'# Perfil de dades — Fase 0 ({dt.date.today().isoformat()})')
    out(f'Fitxer: `{path}`')
    if missing:
        out(f'\n**Pestanyes que falten a l\'export**: {", ".join(missing)}')

    section('1. Files per pestanya')
    table([(TABS[k][0], len(v)) for k, v in d.items()], ['Pestanya', 'Files'])
    extres = [n for n in names if n not in [t[0] for t in TABS.values()]]
    if extres:
        out('Pestanyes no usades per l\'app (o antigues): ' + ', '.join(extres))

    mv = d['moviments']
    section('2. Moviments')
    if not mv:
        out('Cap moviment.')
    else:
        ids = Counter(m.get('id', '') for m in mv)
        dups = {k: c for k, c in ids.items() if c > 1}
        out(f'- Moviments: **{len(mv)}**. Ids duplicats: **{len(dups)}**' + (f' (p. ex. {list(dups)[:5]})' if dups else '') + '.')
        sense_id = sum(1 for m in mv if not m.get('id'))
        out(f'- Moviments sense id: {sense_id}.')
        fmt = Counter()
        for m in mv:
            i = m.get('id', '')
            if re.match(r'^m\d+_[a-z0-9]{4}$', i):
                fmt['m<ms>_<4> (genId)'] += 1
            elif re.match(r'^m\d+_[a-z0-9]{4}0\.\d+', i):
                fmt['genId + Math.random (conté punt)'] += 1
            elif re.match(r'^m\d{12,14}$', i):
                fmt['m<ms> (tancament)'] += 1
            else:
                fmt['altres'] += 1
        out('- Formats d\'id: ' + ', '.join(f'{k}: {v}' for k, v in fmt.items()) + '.')

        bad_data = [m for m in mv if not RE_DATA.match(m.get('data', ''))]
        iso = [m for m in bad_data if RE_ISO.match(m.get('data', ''))]
        out(f'- Dates que no són DD/MM/AAAA: **{len(bad_data)}** (ISO: {len(iso)}, buides/altres: {len(bad_data) - len(iso)}).')
        bad_imp = [m for m in mv if num(m.get('import')) is None]
        neg = [m for m in mv if (num(m.get('import')) or 0) < 0]
        zero = [m for m in mv if num(m.get('import')) == 0]
        out(f'- Imports no numèrics: **{len(bad_imp)}**; negatius: **{len(neg)}**; zero: **{len(zero)}**.')
        sense_any = sum(1 for m in mv if not m.get('anyFiscal'))
        any_rar = Counter(m.get('anyFiscal', '') for m in mv if m.get('anyFiscal') and not RE_ANY.match(m['anyFiscal']))
        out(f'- `anyFiscal` buit: **{sense_any}**; amb format no XX/YY: {dict(any_rar) or 0}.')
        for col in ['tipus', 'via', 'origen', 'origenFons', 'fontImport', 'conciliacio']:
            c = Counter(m.get(col, '') or '(buit)' for m in mv)
            out(f'- `{col}`: ' + ', '.join(f'{k}: {v}' for k, v in c.most_common(12)))
        out()
        out('Camps buits (de més a menys):')
        table(sorted(empties(mv, TABS['moviments'][1]), key=lambda x: -x[1]), ['Camp', 'Buits', '%'])

        h = Counter(m.get('bankHash', '') for m in mv if m.get('bankHash'))
        hd = {k: c for k, c in h.items() if c > 1}
        sense_hash = sum(1 for m in mv if m.get('origen') == 'CSV' and not m.get('bankHash'))
        out(f'- `bankHash` repetits: **{len(hd)}** hashos ({sum(hd.values())} moviments). CSV sense hash: {sense_hash}.')
        if hd:
            ex = list(hd.items())[:5]
            out('  Exemples: ' + '; '.join(f'`{k}` ×{c}' for k, c in ex))
            # repetits amb sentit diferent
            sent = defaultdict(set)
            for m in mv:
                if m.get('bankHash') in hd:
                    sent[m['bankHash']].add(m.get('tipus'))
            mixt = [k for k, v in sent.items() if len(v) > 1]
            out(f'  Dels quals amb ingrés i pagament barrejats: {len(mixt)}.')

        out()
        out('Sumes per any fiscal × tipus (base de la validació):')
        suma = defaultdict(float)
        cnt = Counter()
        for m in mv:
            k = (m.get('anyFiscal') or '(buit)', m.get('tipus') or '(buit)')
            suma[k] += num(m.get('import')) or 0
            cnt[k] += 1
        table([(a, t, cnt[(a, t)], f'{suma[(a, t)]:.2f}') for (a, t) in sorted(suma)], ['Any', 'Tipus', 'Files', 'Suma'])
        out('Sumes per any fiscal × via × tipus:')
        sv = defaultdict(float)
        for m in mv:
            sv[(m.get('anyFiscal') or '(buit)', m.get('via') or '(buit)', m.get('tipus') or '(buit)')] += num(m.get('import')) or 0
        table([(a, v, t, f'{x:.2f}') for (a, v, t), x in sorted(sv.items())], ['Any', 'Via', 'Tipus', 'Suma'])

    section('3. Referències')
    arees = {a.get('id') for a in d['arees']}
    delegs = {x.get('id') for x in d['delegacions']}
    events = {e.get('id') for e in d['events']}
    bancs = {b.get('id') for b in d['bancs']}
    codis = {b.get('codi') for b in d['bancs']}
    usuaris = {u.get('id') for u in d['usuaris']}
    emails = {(u.get('email') or '').lower() for u in d['usuaris']}
    movids = {m.get('id') for m in mv}

    def orfes(recs, col, valid, label, allow_empty=True, allow=()):
        bad = [r for r in recs if (r.get(col) or '') not in valid and r.get(col) not in allow and (r.get(col) or not allow_empty)]
        vals = Counter(r.get(col) for r in bad)
        out(f'- {label}: **{len(bad)}**' + (f' — valors: {dict(vals.most_common(6))}' if bad else ''))

    orfes(mv, 'areaId', arees, 'Moviments amb `areaId` inexistent')
    orfes(mv, 'delegacioId', delegs, 'Moviments amb `delegacioId` inexistent')
    orfes(mv, 'eventId', events, 'Moviments amb `eventId` inexistent')
    orfes(mv, 'via', codis | {'Efectiu', 'TPV', 'app', 'Virtual'}, 'Moviments amb `via` que no és cap banc configurat (ni Efectiu/TPV/app/Virtual)', allow_empty=False)
    sense_area = sum(1 for m in mv if not m.get('areaId'))
    out(f'- Moviments sense àrea: {sense_area}.')
    # delegació incoherent amb àrea
    deleg_area = {x.get('id'): x.get('areaId') for x in d['delegacions']}
    inc = [m for m in mv if m.get('delegacioId') in deleg_area and m.get('areaId') and deleg_area[m['delegacioId']] != m['areaId']]
    out(f'- Moviments amb delegació d\'una altra àrea: {len(inc)}.')
    nomdesf = [m for m in mv if m.get('areaId') in arees and m.get('areaNom') and m['areaNom'] != next((a.get('nom') for a in d['arees'] if a.get('id') == m['areaId']), m['areaNom'])]
    out(f'- Moviments amb `areaNom` desfasat respecte a la taula d\'àrees: {len(nomdesf)}.')

    fons = d['fons']
    out(f'- Fons: {len(fons)}; tipus: ' + ', '.join(f'{k}: {v}' for k, v in Counter(f.get('tipus') for f in fons).most_common()))
    fid = Counter(f.get('id') for f in fons)
    out(f'- Fons amb id duplicat: {sum(1 for c in fid.values() if c > 1)}.')
    f_orf = [f for f in fons if f.get('movimentId') and f['movimentId'] not in movids]
    out(f'- **Fons amb `movimentId` que no existeix (A-24)**: {len(f_orf)}' + (f' — tipus: {dict(Counter(f.get("tipus") for f in f_orf))}' if f_orf else '') + '.')
    f_usr = [f for f in fons if not str(f.get('delegId', '')).startswith('_') and f.get('delegId') not in usuaris and (f.get('delegEmail') or '').lower() not in emails]
    out(f'- Fons d\'un delegat que ja no és usuari: {len(f_usr)}.')
    f_ev = Counter(f.get('delegId') for f in fons if str(f.get('delegId', '')).startswith('_'))
    out(f'- Fons d\'event (`_event`): {dict(f_ev)}.')
    f_bad = [f for f in fons if num(f.get('import')) is None]
    out(f'- Fons amb import no numèric: {len(f_bad)}.')
    reg_nom = sum(1 for f in fons if f.get('registratPer') and '@' not in f['registratPer'])
    out(f'- Fons on `registratPer` és un nom i no un email (A-07): {reg_nom} de {len(fons)}.')

    tc = d['tancaments']
    sint = [t for t in tc if str(t.get('eventId', '')).startswith('_se_')]
    t_orf = [t for t in tc if t.get('eventId') not in events and not str(t.get('eventId', '')).startswith('_se_')]
    out(f'- Tancaments: {len(tc)} (estat: {dict(Counter(t.get("estat") for t in tc))}); amb event sintètic `_se_`: {len(sint)}; amb event inexistent: {len(t_orf)}.')
    pp = d['pres_plans']
    clau = Counter((p.get('any'), p.get('delegacioId'), p.get('eventId'), p.get('concepte')) for p in pp)
    out(f'- Línies de pressupost: {len(pp)}; amb clau natural repetida: {sum(1 for c in clau.values() if c > 1)}; '
        f'`Varis` negatius: {sum(1 for p in pp if p.get("concepte") == "Varis" and (num(p.get("importPla")) or 0) < 0)}; '
        f'amb delegació inexistent: {sum(1 for p in pp if p.get("delegacioId") not in delegs)}.')
    asg = d['assignacions']
    out(f'- Assignacions de directiva: {len(asg)}; amb usuari inexistent: {sum(1 for a in asg if a.get("userId") not in usuaris)}; '
        f'amb àrea/delegació inexistent: {sum(1 for a in asg if a.get("id") not in (arees | delegs))}; '
        f'sense anyFiscal: {sum(1 for a in asg if not a.get("anyFiscal"))}.')

    section('4. Saldos inicials (`15 Saldos`) — A-03')
    sal = d['saldos']
    cont = [s for s in sal if s.get('bancId') not in bancs or not RE_ANY.match(s.get('anyFiscal', ''))]
    out(f'- Files: {len(sal)}; **amb `bancId` desconegut o `anyFiscal` no XX/YY**: {len(cont)}.')
    for s in cont[:10]:
        out(f'  - bancId=`{s.get("bancId")}` anyFiscal=`{s.get("anyFiscal")}` saldoInicial={s.get("saldoInicial")}')
    out()
    out('Bancs configurats: ' + ', '.join(f'{b.get("id")} = {b.get("codi")} ({b.get("tipus")})' for b in d['bancs']))

    section('5. Configuració (`00 Config`) — A-01, A-02')
    cfg = {c.get('clau'): c.get('valor') for c in d['config']}
    out(f'- Claus: {sorted(cfg)}')
    out(f'- `anyFiscal` actiu: `{cfg.get("anyFiscal")}`; `exercicis`: `{cfg.get("exercicis")}`')
    out(f'- `estats_exercicis`: `{cfg.get("estats_exercicis")}`')
    out(f'- `justificant`: `{cfg.get("justificant")}`; `drive_folders`: {"configurades" if cfg.get("drive_folders") not in (None, "", "{}") else "buides"}')
    usr = d['usuaris']
    out(f'- Usuaris: {len(usr)}; rols: {dict(Counter(u.get("rol") for u in usr))}; sense email: {sum(1 for u in usr if not u.get("email"))}; '
        f'emails repetits: {sum(1 for c in Counter((u.get("email") or "").lower() for u in usr).values() if c > 1)}.')
    rols_def = {r.get('id') for r in d['rols']}
    out(f'- Usuaris amb rol que no és a `14 Rols`: {sum(1 for u in usr if u.get("rol") not in rols_def and rols_def)}.')

    section('6. Justificants')
    tk = [m for m in mv if m.get('ticketUrl')]
    drive = [m for m in tk if 'drive.google.com' in m['ticketUrl']]
    out(f'- Moviments amb `ticketUrl`: {len(tk)} (de Drive: {len(drive)}). Despeses manuals sense justificant: '
        f'{sum(1 for m in mv if m.get("origen") == "Manual" and m.get("tipus") == "gast" and not m.get("ticketUrl"))}.')

    nom = re.sub(r'\.xlsx$', '', path, flags=re.I)
    dest = f'{nom}_perfil_dades_{dt.date.today().isoformat()}.md'
    with open(dest, 'w', encoding='utf-8') as f:
        f.write('\n'.join(OUT))
    print(f'\nInforme desat a: {dest}')


if __name__ == '__main__':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass
    if len(sys.argv) != 2:
        print(__doc__)
        sys.exit(1)
    main(sys.argv[1])
